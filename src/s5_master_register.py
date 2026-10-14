"""Step 5: the master register. One Excel file with every datasheet, its metadata and its long-term link.

Sheets:
  Datasheets         one row per datasheet (SharePoint is the source of the metadata)
  SharePoint twins   older SharePoint copies of the same datasheet (the '½' vs '%C2%BD' pairs), not synced
  DAM only           DAM files in the datasheet folders with no SharePoint source (left untouched)
  About              when and from what the file was built

Reads the step 3 plan (data/output/s3) and, if present, the step 4 ledger (ids and links created by a sync).
For the final register after a sync, re-run steps 2 and 3 first so every value is read back from the DAM:
  python src/s2_dam_inventory.py && python src/s3_compare.py && python src/s5_master_register.py
Writes data/output/s5/datasheet_register.xlsx and datasheet_register.csv (same Datasheets sheet, UTF-8).
"""
import csv
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import env  # noqa: E402
from s4_sync import blocked_items, is_blocked, load_ledger  # noqa: E402

S3, S4, OUT = env.output_dir("s3"), env.output_dir("s4"), env.output_dir("s5")

COLUMNS = [  # (header, width, value from plan row p, ledger entry e)
    ("Datasheet", 42, lambda p, e, m: p["dam_name"] or p["file_name"].replace("%C2%BD", "½")),
    ("Long-term link", 70, lambda p, e, m: e.get("link") or p["long_term_link"]),
    ("Item code", 14, lambda p, e, m: m.get("dsitemcode", "")),
    ("Item description", 40, lambda p, e, m: m.get("dsitemdescription", "")),
    ("Item no", 10, lambda p, e, m: m.get("dsitemno", "")),
    ("Language code", 10, lambda p, e, m: m.get("dslcode", "")),
    ("Language", 14, lambda p, e, m: m.get("dslanguage", "")),
    ("LMLID", 8, lambda p, e, m: m.get("dslmlid", "")),
    ("Type", 8, lambda p, e, m: m.get("dstypecode", "")),
    ("Template", 10, lambda p, e, m: m.get("dstemplate", "")),
    ("Page format", 10, lambda p, e, m: m.get("dspageformat", "")),
    ("Proofread", 10, lambda p, e, m: m.get("dsproofread", "")),
    ("Product group", 22, lambda p, e, m: m.get("dsproductgroup", "")),
    ("Sub product group", 22, lambda p, e, m: m.get("dssubproductgroup", "")),
    ("DAM folder", 30, lambda p, e, m: p["dam_folder"] or (p["target_folder"] if e.get("dam_id") else "")),
    ("Target folder", 30, lambda p, e, m: p["target_folder"]),
    ("Created (SharePoint)", 21, lambda p, e, m: m.get("dscreated", "")),
    ("Modified (SharePoint)", 21, lambda p, e, m: m.get("dsmodified", "")),
    ("Size (bytes)", 12, lambda p, e, m: p["spo_size"]),
    ("SharePoint id", 12, lambda p, e, m: m.get("dsid", "")),
    ("SharePoint file name", 42, lambda p, e, m: p["file_name"]),
    ("DAM file id", 38, lambda p, e, m: e.get("dam_id") or p["dam_id"]),
    ("Status", 14, None),
    ("Open actions", 40, None),
]


def open_actions(p: dict, e: dict) -> list:
    done = set(e.get("done", []))
    todo = []
    if p["content"] == "UPLOAD" and "upload" not in done:
        todo.append("upload")
    if p["content"] == "REPLACE" and "replace" not in done:
        todo.append("replace content")
    if (p["metadata_to_set"] or p.get("dup_metadata_keys")) and "metadata" not in done:
        todo.append("update metadata")
    if p["link"] == "CREATE_LINK" and "links" not in done and not e.get("link"):
        todo.append("create link")
    if p["placement"] == "MOVE":
        todo.append(f"move to {p['target_folder']}")
    if p["placement"] == "KEEP_PROTECTED":
        todo.append(f"kept in {p['dam_folder']} (linked)")
    return todo


def write_sheet(wb, title, headers, widths, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, w in enumerate(widths):
        ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = w
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "B2"
    if rows:
        ref = f"A1:{ws.cell(1, len(headers)).column_letter}{len(rows) + 1}"
        tab = Table(displayName=title.replace(" ", "_"), ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(tab)
    return ws


def main():
    plan = [json.loads(l) for l in (S3 / "plan.jsonl").open(encoding="utf-8")]
    ledger = load_ledger()  # ledger.json plus ledger_<name>.json from parallel runs
    comparison = list(csv.DictReader((S3 / "comparison.csv").open(encoding="utf-8-sig")))

    headers = [c[0] for c in COLUMNS]
    blocked = blocked_items()
    rows = []
    for p in sorted(plan, key=lambda p: (p["expected_metadata"].get("dsitemcode", ""), p["file_name"])):
        e = ledger.get(p["file_name"], {})
        m = p["expected_metadata"]
        todo = open_actions(p, e)
        status = "in sync" if not [t for t in todo if not t.startswith("kept in")] else (
            "not in DAM" if not (e.get("dam_id") or p["dam_id"]) else "needs update")
        if is_blocked(p, blocked):
            status, todo = "excluded", ["blocked in config/skip_items.csv, not synced"]
        rows.append([fn(p, e, m) if fn else None for _, _, fn in COLUMNS[:-2]] + [status, "; ".join(todo)])

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "Datasheets", headers, [c[1] for c in COLUMNS], rows)

    twins = [r for r in comparison if r["content"] == "SKIP_OLDER_TWIN"]
    write_sheet(wb, "SharePoint twins", ["Older SharePoint file (not synced)", "Synced from", "Size (bytes)"],
                [48, 48, 12], [[r["file_name"], r["spo_older_twin"], int(r["spo_size"] or 0)] for r in twins])

    orphans = [r for r in comparison if r["content"] == "ORPHAN"]
    write_sheet(wb, "DAM only", ["DAM file", "DAM folder", "Long-term link", "Origin", "Size (bytes)", "DAM file id"],
                [48, 36, 70, 22, 12, 38],
                [[r["file_name"], r["dam_folder"], r["long_term_link"], r["origin"], int(r["dam_size"] or 0), r["dam_id"]]
                 for r in orphans])

    stamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    counts = {}
    for r in rows:
        counts[r[-2]] = counts.get(r[-2], 0) + 1
    about = [["Built", stamp],
             ["Source", "SharePoint sites/Datasheets, library PDF (metadata); AMMshare DAM (folders, ids, links)"],
             ["Datasheets", len(rows)],
             ["With long-term link", sum(1 for r in rows if r[1])],
             *[[f"Status: {k}", v] for k, v in sorted(counts.items())],
             ["SharePoint twins (not synced)", len(twins)],
             ["DAM only (untouched)", len(orphans)],
             ["Sync ledger used", "yes" if ledger else "no (no sync has run yet)"],
             ["Note", "Dates are as stored in SharePoint (UTC). Long-term links never change when a datasheet is updated."]]
    write_sheet(wb, "About", ["Item", "Value"], [32, 100], about)

    xlsx = OUT / "datasheet_register.xlsx"
    tmp = xlsx.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(xlsx)
    with (OUT / "datasheet_register.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(headers)
        w.writerows(rows)
    print(f"{len(rows):,} datasheets ({sum(1 for r in rows if r[1]):,} with link), "
          f"{len(twins):,} twins, {len(orphans):,} DAM-only -> {xlsx}")
    for k, v in sorted(counts.items()):
        print(f"  {k}: {v:,}")


if __name__ == "__main__":
    main()
